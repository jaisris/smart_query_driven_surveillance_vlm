# -*- coding: utf-8 -*-
"""Build the 15-minute mid-term progress deck (.pptx) with speaker notes."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x14, 0x3A, 0x66)
ACCENT = RGBColor(0x1A, 0x73, 0xE8)
GRAY = RGBColor(0x3C, 0x40, 0x43)
GREEN = RGBColor(0x18, 0x80, 0x38)
ORANGE = RGBColor(0xE3, 0x74, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BASE = r"C:\Users\MSUSERSL123\Jaisri\BITS\smart_query_driven_surveillance_vlm"
ARCH = os.path.join(BASE, "Docs", "arch_diagram.png")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_base(title):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.3), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = title
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = NAVY
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.2), Inches(3.2), Pt(3))
    ln.fill.solid()
    ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background()
    return s


def bullets(slide, items, left=0.7, top=1.5, width=12.0, height=5.6, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        if isinstance(it, dict):
            text, lvl = it["t"], it.get("lvl", 0)
            col, sz, mark = it.get("c", GRAY), it.get("s", size), it.get("mark", "•")
        else:
            text, lvl, col, sz, mark = it, 0, GRAY, size, "•"
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pad = "   " if lvl == 0 else "         "
        r = p.add_run()
        r.text = f"{pad}{mark}  {text}"
        r.font.size = Pt(sz)
        r.font.color.rgb = col
        p.space_after = Pt(9)
    return tb


def table(slide, headers, rows, left, top, width, height, col_w=None, fs=15):
    nr, nc = len(rows) + 1, len(headers)
    gt = slide.shapes.add_table(nr, nc, Inches(left), Inches(top), Inches(width), Inches(height)).table
    for j, h in enumerate(headers):
        c = gt.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        for para in c.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                run.font.size = Pt(fs)
                run.font.bold = True
                run.font.color.rgb = WHITE
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.text = str(val)
            for para in c.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(fs)
                    run.font.color.rgb = GRAY
    if col_w:
        for j, w in enumerate(col_w):
            gt.columns[j].width = Inches(w)
    return gt


# ---------------- SLIDE 1: Title ----------------
s = prs.slides.add_slide(BLANK)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.35), Inches(13.333), Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
tb = s.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(1.5))
tf = tb.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run()
r.text = ("Vision-Language Based Smart Surveillance System for\n"
          "Query-Driven Video Retrieval and Suspicious Activity Detection")
r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
sub = s.shapes.add_textbox(Inches(0.8), Inches(2.7), Inches(11.7), Inches(3.2))
stf = sub.text_frame; stf.word_wrap = True
lines = [
    ("Mid-Term Progress Review  ·  Dissertation (S2-25_DISSERTATION-NSP4)", 18, True, NAVY),
    ("Jaisri S   ·   2024AA05138   ·   M.Tech AI & ML", 16, False, GRAY),
    ("Supervisor: Savitha C, IBM India Pvt. Ltd., Bangalore", 16, False, GRAY),
    ("BITS Pilani  ·  Work Integrated Learning Programmes  ·  July 2026", 14, False, GRAY),
]
for i, (t, sz, b, col) in enumerate(lines):
    p = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    rr = p.add_run(); rr.text = t
    rr.font.size = Pt(sz); rr.font.bold = b; rr.font.color.rgb = col
    p.space_after = Pt(10)
notes(s, "(~0:30) Good morning, and thank you for your time. I'm Jaisri, presenting the mid-term "
      "progress of my dissertation: a vision-language surveillance system that lets an operator search "
      "recorded video with plain-language queries and, at the same time, flags suspicious activity. In "
      "the next fifteen minutes I'll cover the problem, the architecture I've built, the results I'm "
      "already getting on real footage, and my plan for the remaining weeks.")

# ---------------- SLIDE 2: Problem ----------------
s = slide_base("The Problem")
bullets(s, [
    "A site may run dozens of cameras, but the footage is raw and unlabelled.",
    "Finding one event after the fact usually means watching the recording end to end.",
    "Cost grows with recording length and depends entirely on operator attention.",
    "Fixed alarm rules only fire on presets, not on how people describe an event.",
])
notes(s, "(~1:30) Surveillance cameras record continuously, but almost nothing they capture is indexed. "
      "When something must be found afterwards, the standard method is still a person scrubbing through "
      "hours of video. That doesn't scale and is error-prone. Rule-based alarms help only when the event "
      "can be reduced to a fixed trigger, like motion in a zone. But an operator thinks in natural "
      "language: 'a person loitering near the gate'. My project bridges that gap, and also watches for "
      "suspicious behaviour automatically.")

# ---------------- SLIDE 3: Objectives ----------------
s = slide_base("Objectives & Scope")
bullets(s, [
    "Retrieve relevant video segments from long footage using natural-language queries.",
    "Detect and track people and vehicles across frames.",
    "Match text to visual content via a shared vision-language embedding.",
    "Return time-stamped segments, not isolated frames.",
    "Flag suspicious activity: loitering and zone intrusion.",
    "Evaluate each stage with standard metrics.",
])
notes(s, "(~1:00) My objectives fall into two halves. First, query-driven retrieval: detect and track "
      "objects, embed both video and query into a shared space, and return usable time segments ranked by "
      "relevance. Second, suspicious-activity detection, specifically loitering and intrusion, in the same "
      "pass. Underpinning both is an evaluation layer. Scope is recorded footage rather than live "
      "streaming, which keeps the problem tractable for a dissertation.")

# ---------------- SLIDE 4: Architecture ----------------
s = slide_base("System Architecture")
bullets(s, [
    {"t": "Linear pipeline: video → detection + tracking → CLIP embedding → FAISS → segments → UI.", "s": 15},
    {"t": "Query branch: text → CLIPEncoder → same index.", "s": 15},
    {"t": "Anomaly branch: track histories → AnomalyEngine → UI.", "s": 15},
    {"t": "Content cache (SHA-256) reuses detections and embeddings.", "s": 15},
], left=0.6, top=1.5, width=4.7, size=15)
if os.path.exists(ARCH):
    s.shapes.add_picture(ARCH, Inches(5.35), Inches(1.45), height=Inches(5.7))
notes(s, "(~1:30) This is the architecture. Video is sampled every fifteenth frame. Each frame goes "
      "through YOLOv8 for detection and DeepSORT for tracking, so objects keep a stable identity. CLIP "
      "embeds each frame into a 512-dimensional vector, indexed in FAISS. The text query is encoded by "
      "the same CLIP model into the same space, so a sentence is compared directly against the imagery. "
      "Matches are merged into time segments; in parallel the track histories feed the anomaly engine. "
      "One key point is the content cache: results are keyed by the video hash, so re-analysing the same "
      "clip is near-instant.")

# ---------------- SLIDE 5: Pipeline ----------------
s = slide_base("How the Pipeline Works")
bullets(s, [
    "VideoLoader samples frames (OpenCV).",
    "FrameProcessor: YOLOv8 detections + DeepSORT tracks (BGR to RGB handled here).",
    "EmbeddingBuilder: CLIP frame vectors (512-d), cached as .npy.",
    "SimilaritySearch: FAISS IndexFlatIP equals cosine similarity.",
    "TemporalLocalizer: merges nearby top-K frames into segments.",
    "AnomalyEngine: rule-based loitering + intrusion.",
])
notes(s, "(~1:30) A little more concretely: the loader produces sampled frames; the frame processor runs "
      "detection and tracking and returns boxes, classes and stable track IDs. The embedding builder "
      "converts frames to CLIP vectors and caches them. At query time, FAISS inner-product search over "
      "normalised vectors is exactly cosine similarity. The temporal localizer groups nearby high-scoring "
      "frames into continuous segments with start and end times. The anomaly engine works off track "
      "histories in parallel. Every module is separate and independently testable.")

# ---------------- SLIDE 6: Modules ----------------
s = slide_base("Modules Implemented")
bullets(s, [
    "Configuration & utilities  ·  Video loading & cache",
    "YOLOv8 detector  ·  DeepSORT tracker",
    "CLIP encoder (image + text)",
    "FAISS retrieval  ·  Temporal localizer",
    "Anomaly engine (loitering, intrusion)",
    "Evaluation modules (4 tasks)  ·  Streamlit UI  ·  unit tests",
])
notes(s, "(~1:00) In terms of what exists in the codebase today, all of these modules are implemented, "
      "not just planned. There is a clean separation between configuration, data handling, the models, "
      "retrieval, anomaly analysis, evaluation and the interface, plus a set of unit tests. So this is a "
      "working, modular prototype rather than a design on paper.")

# ---------------- SLIDE 7: Tech stack ----------------
s = slide_base("Technology Stack")
bullets(s, [
    {"t": "Object detection: YOLOv8 nano (Ultralytics), real-time and COCO-pretrained.", "s": 16},
    {"t": "Restricted to person and vehicle classes; confidence threshold 0.40.", "lvl": 1, "s": 15},
    {"t": "Multi-object tracking: DeepSORT, Kalman motion plus deep appearance features.", "s": 16},
    {"t": "Keeps stable identities through short occlusions.", "lvl": 1, "s": 15},
    {"t": "Vision-language: CLIP ViT-B/32 (HuggingFace), frames and text in one 512-d space.", "s": 16},
    {"t": "Retrieval: FAISS IndexFlatIP, exact cosine search on normalised vectors.", "s": 16},
    {"t": "Video I/O: OpenCV;  UI: Streamlit;  Evaluation: motmetrics, scikit-learn.", "s": 16},
    {"t": "Config-driven: frame_skip = 15, top_k = 10 and thresholds, all in one YAML file.", "s": 16},
], size=16)
notes(s, "(~0:45) The stack is mainstream and reproducible: YOLOv8, DeepSORT, CLIP ViT-B/32 through "
      "HuggingFace, FAISS for search, Streamlit for the demo. Evaluation uses motmetrics and scikit-learn, "
      "the standard libraries for these metrics. Everything tunable lives in a single config file.")

# ---------------- SLIDE 8: Detection/Tracking results ----------------
s = slide_base("Results: Detection & Tracking")
bullets(s, [
    {"t": "Test footage: 640 x 360, 30 fps, 58.4 s, 1751 frames.", "s": 16},
    {"t": "Sampled every 15th frame, giving 117 frames analysed.", "lvl": 1, "s": 15},
    {"t": "117 unique object tracks maintained across the clip.", "c": NAVY, "s": 16},
    {"t": "Mean track length 44 frames; longest track 102 frames.", "lvl": 1, "s": 15},
    {"t": "Class mix: car 108, bus 5, person 3, truck 1, a vehicle-heavy scene.", "s": 16},
    {"t": "Data-quality guard: 2 blank frames detected and removed, 115 clean frames indexed.", "c": ORANGE, "s": 16},
    {"t": "Speed: cold run about 132 s on CPU; cached re-run about 2.7 s.", "s": 16},
    {"t": "All figures taken directly from the system's own run logs.", "s": 15},
], size=16)
notes(s, "(~1:30) Results on real footage. I ran the full pipeline on a 58-second traffic clip. From the "
      "117 sampled frames it produced 117 object tracks, with a clear breakdown: 108 cars, five buses, "
      "three people, one truck. That breakdown is itself a correctness signal, it recognises this as a "
      "vehicle-heavy road scene. On CPU the cold run takes about two minutes, dominated by detection and "
      "tracking, but the cache brings a repeat run under three seconds. All figures are from the system's "
      "own logs, not estimates.")

# ---------------- SLIDE 9: Retrieval results (table) ----------------
s = slide_base("Results: Query-Driven Retrieval")
table(s, ["Query", "Top cosine", "Peak at", "Segments"], [
    ["\"buses passing\"", "0.306", "55.5 s", "2"],
    ["\"person walking\"", "0.246", "17.0 s", "3"],
    ["\"cars colliding\"", "0.310", "11.0 s", "4"],
], left=1.0, top=1.5, width=9.3, height=1.85, col_w=[4.1, 1.9, 1.7, 1.6], fs=15)
bullets(s, [
    {"t": "Three natural-language queries against the indexed frames (top-10, cosine).", "s": 15},
    {"t": "\"buses passing\" peaks at 55.5 s, exactly where the buses appear.", "c": NAVY, "s": 15},
    {"t": "\"cars colliding\" (0.31) matches the densest multi-car frames; no real collision in this clip.", "c": ORANGE, "s": 15},
    {"t": "Scores 0.23 to 0.31 are the normal range for raw CLIP ViT-B/32 similarity.", "s": 15},
    {"t": "Nearby hits merged into segments (2 s gap); warm query latency 0.1 to 0.4 s.", "s": 15},
], top=3.75, size=15)
notes(s, "(~1:30) This is retrieval working. I issued three natural-language queries. 'buses passing' peaks "
      "at 55.5 seconds, exactly where the buses appear in the clip, and 'person walking' returns pedestrian "
      "moments. 'cars colliding' scores highest at about 0.31, but I want to be honest here: there is no "
      "actual collision in this footage, so what CLIP matches is the densest multi-car frames that visually "
      "resemble the query. That is an important nuance, the retrieval finds visually similar moments, and "
      "verifying a true collision would need the dedicated anomaly path. Cosine scores of 0.23 to 0.31 are "
      "normal for CLIP ViT-B/32, and each warm query returns in a fraction of a second.")

# ---------------- SLIDE 10: Anomaly results ----------------
s = slide_base("Results: Suspicious Activity Detection")
bullets(s, [
    {"t": "Two rule-based detectors run on the DeepSORT track histories:", "s": 16},
    {"t": "Loitering: a person staying within an 80 px radius for at least 30 s.", "lvl": 1, "s": 15},
    {"t": "Intrusion: a track entering a user-defined restricted zone (ROI polygon).", "lvl": 1, "s": 15},
    {"t": "Severity = observed dwell time divided by the threshold.", "s": 16},
    {"t": "Traffic clip: 0 events, correct (only 3 short person tracks, none dwell 30 s).", "c": NAVY, "s": 16},
    {"t": "Pedestrian clip (351 frames, 135 tracks): 17 loitering events, severity 1.02 to 5.18.", "c": ORANGE, "s": 16},
    {"t": "Person-only filter avoids false alarms from parked or slow vehicles.", "s": 15},
    {"t": "Extensible: optional VadCLIP path for learned anomaly scoring.", "s": 15},
], size=16)
notes(s, "(~1:15) For anomaly detection I want to be precise. On the traffic clip, zero events were "
      "raised, and that is the correct outcome, not a failure. Loitering requires a person to stay within "
      "a small radius for at least thirty seconds, and this clip has only three brief person tracks among "
      "mostly vehicles. To show the detector works, I ran a separate pedestrian clip: there it produced 17 "
      "loitering events, with severity from just over the threshold up to about five times it. So the "
      "logic is sound and sensitive to real dwell behaviour.")

# ---------------- SLIDE 11: Evaluation ----------------
s = slide_base("Evaluation Methodology")
bullets(s, [
    {"t": "Detection, mAP@0.5: box precision vs ground truth via IoU matching (11-point AP).", "s": 16},
    {"t": "Tracking, MOTA: overall errors, misses, false positives and ID switches.", "s": 16},
    {"t": "Tracking, IDF1: identity consistency across frames; plus MT and ML.", "lvl": 1, "s": 15},
    {"t": "Retrieval: Precision@K, Recall@K, NDCG@K (rank quality) and MRR (first hit).", "s": 16},
    {"t": "Anomaly: AUC-ROC (separability), Average Precision and EER.", "s": 16},
    {"t": "All four metric modules implemented and unit-tested against known inputs.", "s": 15},
    {"t": "Datasets: UCF-Crime (anomaly), MOT17 (tracking), COCO (detection).", "s": 15},
    {"t": "Benchmark scores deferred until run on labelled data; real outputs shown instead.", "c": ORANGE, "s": 15},
], size=16)
notes(s, "(~1:15) For each stage I implemented the standard metric: mAP for detection, MOTA and IDF1 for "
      "tracking, ranking metrics for retrieval, and frame-level AUC-ROC for anomaly detection. The code is "
      "written and unit-tested. What I am deliberately not doing is quoting benchmark numbers yet, because "
      "those require labelled ground truth from datasets like UCF-Crime and MOT17, and running that at "
      "scale is my next phase. I would rather show genuine system outputs than fabricated scores.")

# ---------------- SLIDE 12: Progress summary ----------------
s = slide_base("Progress Summary")
bullets(s, [
    {"t": "Core pipeline: sampling, YOLOv8, DeepSORT, CLIP, FAISS, temporal segments.", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Query-driven retrieval working and validated on real footage.", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Rule-based anomaly detection (loitering, intrusion) firing correctly.", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Streamlit demo: upload, query, suggestions, segment playback, anomaly panel.", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Evaluation harness and unit tests; content-hash caching for fast re-runs.", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Robustness: real defects found and fixed in testing (blank frames, UI callback).", "mark": "✓", "c": GREEN, "s": 17},
    {"t": "Benchmarking on labelled datasets: in progress.", "mark": "→", "c": ORANGE, "s": 17},
], size=17)
notes(s, "(~0:45) To summarise: the entire core pipeline is complete and validated end to end on real "
      "video. Retrieval works, anomaly detection works, there is a working demo, and the evaluation "
      "harness is ready. The project has clearly moved past design into a functioning prototype. The one "
      "thing still in progress is the large-scale quantitative benchmarking.")

# ---------------- SLIDE 13: Future plan (table) ----------------
s = slide_base("Future Plan  (Remaining ~4 Weeks)")
table(s, ["Work", "Timeline"], [
    ["Full UCF-Crime evaluation & benchmarking (mAP, MOTA/IDF1, AUC-ROC)", "Now – 15 Jul"],
    ["Code optimisation (GPU, batching) & alternative VLMs (ViT-L/14, BLIP-2, SigLIP)", "10 – 20 Jul"],
    ["Robustness study (crowding, occlusion, lighting) & anomaly tuning", "18 – 26 Jul"],
    ["Dissertation writing & final submission", "27 Jul – Aug"],
], left=0.8, top=1.55, width=11.7, height=2.5, col_w=[8.4, 3.3], fs=15)
bullets(s, [
    {"t": "Priority: obtain quantitative benchmark scores on labelled datasets.", "s": 15},
    {"t": "Optimisation matters: UCF-Crime videos are far larger than the demo clips.", "s": 15},
    {"t": "Explore stronger VLMs (ViT-L/14, BLIP-2, SigLIP) for subtle, less literal queries.", "s": 15},
], top=4.35, size=15)
notes(s, "(~1:00) For the remaining four weeks, my priority is running the full evaluation on UCF-Crime "
      "and MOT17 to get benchmark numbers. In parallel I'll optimise the pipeline, mainly GPU and "
      "batching, because UCF-Crime videos are far larger than the demo clips, and trial a stronger "
      "vision-language model such as ViT-L/14 or BLIP-2. Then a robustness study and threshold tuning, and "
      "finally the write-up for submission in August. This keeps me aligned with the plan of work in my "
      "abstract.")

# ---------------- SLIDE 14: Closing ----------------
s = slide_base("Closing & Questions")
bullets(s, [
    "Delivered: natural-language queries return ranked, time-stamped video segments.",
    "One pipeline combining object detection, tracking and vision-language matching.",
    "Context-aware suspicious-activity detection: loitering and intrusion.",
    "Validated on real footage with genuine, logged results, not simulated.",
    "Real data-quality and UI defects found and fixed during testing.",
    "Clear, dated plan through benchmarking to the August submission.",
    {"t": "Thank you.  Questions welcome.", "c": NAVY, "s": 20},
], size=17)
notes(s, "(~0:30) In closing: I have a working, modular system that turns a natural-language query into "
      "ranked, time-stamped video segments and flags suspicious activity, validated on real footage with "
      "measurable results. The path to submission is well defined. Thank you, and I'd be glad to take your "
      "questions.")

import time as _time
_names = ["Midterm_Presentation_2024AA05138.pptx",
          "Midterm_Presentation_2024AA05138_v2.pptx",
          "Midterm_Presentation_2024AA05138_v3.pptx",
          f"Midterm_Presentation_2024AA05138_{int(_time.time())}.pptx"]
out = None
for _n in _names:
    try:
        out = os.path.join(BASE, "Docs", _n)
        prs.save(out)
        break
    except PermissionError:
        out = None
        continue
if out is None:
    raise SystemExit("All candidate files are locked; close PowerPoint and retry.")
if not out.endswith("2024AA05138.pptx"):
    print("NOTE: the canonical file was open/locked; saved to a fallback name instead.")
print("SAVED ->", out)
print("slides:", len(prs.slides._sldIdLst))
